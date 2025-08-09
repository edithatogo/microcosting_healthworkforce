from pathlib import Path

import openpyxl


def inspect_excel_formulas_detailed() -> None:
    """
    Loads the Excel workbook and inspects the formulas for specific rows,
    including both Base Salary and Base Salary + Costs columns.
    """
    file_path = Path("data") / "Labour Model - v3, 20240328.xlsx"
    personnel_types_to_check = {
        "Lab technician",
        "Medical scientist",
        "Admin staff",
        "Bioinformatics staff",
        "Pathologists",
    }

    try:
        workbook = openpyxl.load_workbook(file_path, data_only=False)
        sheet = workbook["Labour - Prices"]

        print("--- Detailed Formula Inspection in 'Labour - Prices' Sheet ---")

        # Columns:
        # B (2): Personnel Type
        # S (19), T (20), U (21): Base Salary (LL, Mid, UL)
        # V (22), W (23), X (24): Base Salary + Costs (LL, Mid, UL)

        found_personnel = set()

        for row_index in range(4, sheet.max_row + 1):
            personnel_type = sheet.cell(row=row_index, column=2).value

            if (
                personnel_type in personnel_types_to_check
                and personnel_type not in found_personnel
            ):
                print(f"\nPersonnel Type: {personnel_type} (Row {row_index})")

                # Base Salary Formulas
                formula_s = sheet.cell(row=row_index, column=19).value
                formula_t = sheet.cell(row=row_index, column=20).value
                formula_u = sheet.cell(row=row_index, column=21).value
                print("  Base Salary Formulas:")
                print(f"    Col S (LL): {formula_s}")
                print(f"    Col T (Mid): {formula_t}")
                print(f"    Col U (UL): {formula_u}")

                # Base Salary + Costs Formulas
                formula_v = sheet.cell(row=row_index, column=22).value
                formula_w = sheet.cell(row=row_index, column=23).value
                formula_x = sheet.cell(row=row_index, column=24).value
                print("  Base Salary + Costs Formulas:")
                print(f"    Col V (LL): {formula_v}")
                print(f"    Col W (Mid): {formula_w}")
                print(f"    Col X (UL): {formula_x}")

                found_personnel.add(personnel_type)

            if len(found_personnel) == len(personnel_types_to_check):
                break

    except FileNotFoundError:
        print(f"Error: File not found at {file_path}")
    except Exception as e:
        print(f"An error occurred: {e}")


if __name__ == "__main__":
    inspect_excel_formulas_detailed()
